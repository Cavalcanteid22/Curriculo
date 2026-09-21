# -*- coding: utf-8 -*-
"""
Gera o capítulo de dissertação que descreve a construção do ELO-SIS.

As listagens de código são extraídas diretamente dos arquivos-fonte, e não
transcritas: qualquer alteração no programa se reflete no capítulo na próxima
geração, o que elimina a divergência entre o que o texto afirma e o que o
produto faz.
"""

from __future__ import annotations

import ast
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
FONTE = "Arial"
FONTE_CODIGO = "Consolas"
COR_TITULO = RGBColor(0x1F, 0x4E, 0x79)
COR_NOTA = RGBColor(0x52, 0x51, 0x4E)
COR_COMENTARIO = RGBColor(0x00, 0x80, 0x00)


# --------------------------------------------------------------------------- #
# Extração de código a partir dos arquivos-fonte
# --------------------------------------------------------------------------- #

def _fonte(modulo: str) -> str:
    return (RAIZ / "elosis" / f"{modulo}.py").read_text(encoding="utf-8")


def extrair_funcao(modulo: str, nome: str, sem_docstring: bool = False,
                   ate_linha: int | None = None) -> str:
    """Devolve o código-fonte de uma função ou classe, tal como está no arquivo."""
    texto = _fonte(modulo)
    arvore = ast.parse(texto)
    linhas = texto.splitlines()
    for no in ast.walk(arvore):
        if isinstance(no, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            if no.name != nome:
                continue
            inicio = no.lineno - 1
            fim = no.end_lineno
            corpo = linhas[inicio:fim]
            if sem_docstring and ast.get_docstring(no):
                # Remove a docstring, preservando a assinatura.
                primeiro = no.body[0]
                corpo = (linhas[inicio:primeiro.lineno - 1]
                         + linhas[primeiro.end_lineno:fim])
            if ate_linha:
                corpo = corpo[:ate_linha] + ["    # (...)"]
            return "\n".join(corpo).rstrip()
    raise ValueError(f"{nome} não encontrado em {modulo}.py")


def extrair_atribuicao(modulo: str, nome: str, ate_linha: int | None = None) -> str:
    """Devolve o código de uma atribuição de módulo (constante ou estrutura)."""
    texto = _fonte(modulo)
    arvore = ast.parse(texto)
    linhas = texto.splitlines()
    for no in arvore.body:
        alvos = []
        if isinstance(no, ast.Assign):
            alvos = [a.id for a in no.targets if isinstance(a, ast.Name)]
        elif isinstance(no, ast.AnnAssign) and isinstance(no.target, ast.Name):
            alvos = [no.target.id]
        if nome in alvos:
            corpo = linhas[no.lineno - 1:no.end_lineno]
            if ate_linha:
                corpo = corpo[:ate_linha] + ["    # (...)"]
            return "\n".join(corpo).rstrip()
    raise ValueError(f"{nome} não encontrado em {modulo}.py")


def extrair_trecho(modulo: str, inicio: str, fim: str | None = None,
                   maximo: int | None = None) -> str:
    """Devolve o trecho entre duas marcas textuais do arquivo."""
    texto = _fonte(modulo)
    posicao = texto.index(inicio)
    recorte = texto[posicao:]
    if fim:
        recorte = recorte[:recorte.index(fim)]
    linhas = recorte.rstrip().splitlines()
    if maximo:
        linhas = linhas[:maximo] + ["    # (...)"]
    return "\n".join(linhas)


# --------------------------------------------------------------------------- #
# Formatação ABNT
# --------------------------------------------------------------------------- #

def configurar(documento: Document) -> None:
    """Margens, fonte e espaçamento conforme a ABNT NBR 14724."""
    for secao in documento.sections:
        secao.top_margin = Cm(3)
        secao.left_margin = Cm(3)
        secao.bottom_margin = Cm(2)
        secao.right_margin = Cm(2)

    normal = documento.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    formato = normal.paragraph_format
    formato.line_spacing = 1.5
    formato.space_after = Pt(0)
    formato.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    formato.first_line_indent = Cm(1.25)

    # O template padrão do python-docx grava <w:zoom> sem o atributo
    # obrigatório w:percent, o que viola o esquema do formato. Word tolera,
    # mas um arquivo destinado a banca não deve carregar erro de validação.
    elemento = documento.settings.element
    for zoom in elemento.findall(qn("w:zoom")):
        if zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")


def titulo(documento: Document, texto: str, nivel: int = 1,
           nova_pagina: bool = False) -> None:
    if nova_pagina:
        documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    paragrafo = documento.add_paragraph(style=documento.styles[f"Heading {min(nivel, 4)}"])
    paragrafo.paragraph_format.first_line_indent = Cm(0)
    paragrafo.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
    paragrafo.paragraph_format.space_after = Pt(12)
    paragrafo.paragraph_format.line_spacing = 1.5
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    execucao = paragrafo.add_run(texto.upper() if nivel == 1 else texto)
    execucao.bold = True
    execucao.font.size = Pt(12)
    execucao.font.name = FONTE
    execucao.font.color.rgb = COR_TITULO


def p(documento: Document, texto: str, recuo: bool = True) -> None:
    paragrafo = documento.add_paragraph()
    paragrafo.paragraph_format.first_line_indent = Cm(1.25 if recuo else 0)
    paragrafo.paragraph_format.space_after = Pt(6)
    paragrafo.paragraph_format.line_spacing = 1.5
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    execucao = paragrafo.add_run(texto)
    execucao.font.size = Pt(12)
    execucao.font.name = FONTE


def citacao(documento: Document, texto: str, fonte: str) -> None:
    """Citação direta com mais de três linhas: recuo de 4 cm, corpo 10, simples."""
    paragrafo = documento.add_paragraph()
    paragrafo.paragraph_format.left_indent = Cm(4)
    paragrafo.paragraph_format.first_line_indent = Cm(0)
    paragrafo.paragraph_format.line_spacing = 1.0
    paragrafo.paragraph_format.space_before = Pt(6)
    paragrafo.paragraph_format.space_after = Pt(6)
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    execucao = paragrafo.add_run(f"{texto} ({fonte})")
    execucao.font.size = Pt(10)
    execucao.font.name = FONTE


def lista(documento: Document, itens: list[str], numerada: bool = False) -> None:
    for posicao, item in enumerate(itens, start=1):
        paragrafo = documento.add_paragraph()
        paragrafo.paragraph_format.left_indent = Cm(1.25)
        paragrafo.paragraph_format.first_line_indent = Cm(-0.5)
        paragrafo.paragraph_format.space_after = Pt(4)
        paragrafo.paragraph_format.line_spacing = 1.5
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        execucao = paragrafo.add_run(("%d. " % posicao if numerada else "• ") + item)
        execucao.font.size = Pt(12)
        execucao.font.name = FONTE


def legenda(documento: Document, texto: str, acima: bool = True) -> None:
    paragrafo = documento.add_paragraph()
    paragrafo.paragraph_format.first_line_indent = Cm(0)
    paragrafo.paragraph_format.space_before = Pt(10 if acima else 2)
    paragrafo.paragraph_format.space_after = Pt(2 if acima else 12)
    paragrafo.paragraph_format.line_spacing = 1.0
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    execucao = paragrafo.add_run(texto)
    execucao.font.size = Pt(10)
    execucao.font.name = FONTE
    execucao.bold = acima
    if not acima:
        execucao.font.color.rgb = COR_NOTA


_RE_COMENTARIO = re.compile(r"(#.*)$")


def programa(documento: Document, numero: int, titulo_texto: str, codigo: str,
             arquivo: str, corpo: int = 8) -> int:
    """Insere uma listagem de código com numeração, moldura e fonte monoespaçada.

    Os comentários recebem cor distinta: em código comentado em português, como
    é o caso, eles carregam parte da justificativa das decisões e precisam ser
    distinguíveis das instruções.
    """
    legenda(documento, f"Programa {numero} — {titulo_texto}", acima=True)

    tabela = documento.add_table(rows=1, cols=1)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    celula = tabela.rows[0].cells[0]
    celula.text = ""
    _sombrear(celula, "F7F7F5")

    for indice, linha in enumerate(codigo.splitlines()):
        paragrafo = celula.paragraphs[0] if indice == 0 else celula.add_paragraph()
        paragrafo.paragraph_format.first_line_indent = Cm(0)
        paragrafo.paragraph_format.space_after = Pt(0)
        paragrafo.paragraph_format.space_before = Pt(0)
        paragrafo.paragraph_format.line_spacing = 1.0
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT

        correspondencia = _RE_COMENTARIO.search(linha)
        if correspondencia and not _dentro_de_texto(linha, correspondencia.start()):
            antes = linha[:correspondencia.start()]
            comentario = correspondencia.group(1)
            if antes:
                execucao = paragrafo.add_run(antes)
                execucao.font.size = Pt(corpo)
                execucao.font.name = FONTE_CODIGO
            execucao = paragrafo.add_run(comentario)
            execucao.font.size = Pt(corpo)
            execucao.font.name = FONTE_CODIGO
            execucao.font.color.rgb = COR_COMENTARIO
            execucao.italic = True
        else:
            execucao = paragrafo.add_run(linha if linha.strip() else " ")
            execucao.font.size = Pt(corpo)
            execucao.font.name = FONTE_CODIGO

    legenda(documento, f"Fonte: elaboração própria. Arquivo {arquivo}.", acima=False)
    return numero + 1


def _dentro_de_texto(linha: str, posicao: int) -> bool:
    """Evita tratar como comentário um '#' que esteja dentro de uma string."""
    trecho = linha[:posicao]
    return (trecho.count('"') % 2 == 1) or (trecho.count("'") % 2 == 1)


def tabela(documento: Document, numero: int, titulo_texto: str,
           colunas: list[str], linhas: list[list], fonte: str,
           larguras: list[float] | None = None, corpo: int = 9) -> int:
    legenda(documento, f"Tabela {numero} — {titulo_texto}", acima=True)
    t = documento.add_table(rows=1, cols=len(colunas))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for indice, nome in enumerate(colunas):
        celula = t.rows[0].cells[indice]
        celula.text = ""
        paragrafo = celula.paragraphs[0]
        paragrafo.paragraph_format.first_line_indent = Cm(0)
        paragrafo.paragraph_format.space_after = Pt(0)
        paragrafo.paragraph_format.line_spacing = 1.0
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        execucao = paragrafo.add_run(str(nome))
        execucao.bold = True
        execucao.font.size = Pt(corpo)
        execucao.font.name = FONTE
        execucao.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(celula, "1F4E79")

    for numero_linha, linha in enumerate(linhas):
        celulas = t.add_row().cells
        for indice, valor in enumerate(linha[:len(colunas)]):
            celula = celulas[indice]
            celula.text = ""
            paragrafo = celula.paragraphs[0]
            paragrafo.paragraph_format.first_line_indent = Cm(0)
            paragrafo.paragraph_format.space_after = Pt(0)
            paragrafo.paragraph_format.line_spacing = 1.0
            paragrafo.alignment = (WD_ALIGN_PARAGRAPH.LEFT if indice == 0
                                   else WD_ALIGN_PARAGRAPH.CENTER)
            execucao = paragrafo.add_run("" if valor is None else str(valor))
            execucao.font.size = Pt(corpo)
            execucao.font.name = FONTE
        if numero_linha % 2 == 1:
            for celula in celulas:
                _sombrear(celula, "F2F2F2")

    if larguras:
        for linha in t.rows:
            for indice, largura in enumerate(larguras[:len(colunas)]):
                linha.cells[indice].width = Cm(largura)

    legenda(documento, fonte, acima=False)
    return numero + 1


def figura(documento: Document, numero: int, titulo_texto: str, caminho: Path,
           fonte: str, largura: float = 15.0) -> int:
    if not Path(caminho).exists():
        return numero
    legenda(documento, f"Figura {numero} — {titulo_texto}", acima=True)
    paragrafo = documento.add_paragraph()
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo.paragraph_format.first_line_indent = Cm(0)
    paragrafo.paragraph_format.space_after = Pt(2)
    paragrafo.add_run().add_picture(str(caminho), width=Cm(largura))
    legenda(documento, fonte, acima=False)
    return numero + 1


def _sombrear(celula, cor: str) -> None:
    elemento = OxmlElement("w:shd")
    elemento.set(qn("w:val"), "clear")
    elemento.set(qn("w:color"), "auto")
    elemento.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(elemento)


def numerar_paginas(documento: Document) -> None:
    rodape = documento.sections[0].footer
    paragrafo = rodape.paragraphs[0]
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    execucao = paragrafo.add_run()
    for etiqueta, valor in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if etiqueta == "instr":
            elemento = OxmlElement("w:instrText")
            elemento.text = valor
        else:
            elemento = OxmlElement("w:fldChar")
            elemento.set(qn("w:fldCharType"), etiqueta)
        execucao._r.append(elemento)
    execucao.font.size = Pt(10)
    execucao.font.name = FONTE
